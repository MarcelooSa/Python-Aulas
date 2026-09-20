import streamlit as st
import time
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# ========== CONSTANTES ==========
MESES = [
    "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"
]

# ========== INICIALIZAÇÃO DO ESTADO ==========
def inicializar_estado():
    """Inicializa todos os estados da sessão"""
    estados_padrao = {
        "itens": [],
        "pagina": "Home",
        "editando": None,
        "aviso_edicao": False,
        "orcamentos_mensais": {},
        "opcoes_cadastro": {},
        "mostrar_tabela": False,
        "mostrar_tabela_historico": False,
    }
    
    for chave, valor in estados_padrao.items():
        if chave not in st.session_state:
            st.session_state[chave] = valor

# ========== CONFIGURAÇÃO DA PÁGINA ==========
st.set_page_config(page_title="SmartList", page_icon="🛒", layout="wide")

st.markdown("""
<style>
    div[data-testid="stSidebarNav"] {
        display: none !important;
    }
</style>
""", unsafe_allow_html=True)

# ========== VERIFICAÇÃO DE LOGIN ==========
if not st.session_state.get("usuario_logado", False):
    st.switch_page("pages/login.py")
    st.stop()

nome_usuario = st.session_state.get("usuario_nome", "Usuário")

# ========== INICIALIZAR ESTADO ==========
inicializar_estado()

# ========== SIDEBAR ==========
with st.sidebar:
    st.title("🛒 SmartList")
    st.write(f"👤 {nome_usuario}")
    st.markdown("---")
    
    if st.button("🏠 Home"):
        st.session_state.pagina = "Home"
        st.rerun()

    if st.button("➕ Cadastro de Itens"):
        st.session_state.pagina = "Cadastro"
        st.rerun()

    if st.button("📋 Listagem de Itens"):
        st.session_state.pagina = "Listagem"
        st.rerun()

    if st.button("💰 Cálculo e Orçamento"):
        st.session_state.pagina = "Orcamento"
        st.rerun()

    if st.button("📊 Histórico Mensal"):
        st.session_state.pagina = "Historico"
        st.rerun()

    if st.button("ℹ️ Sobre"):
        st.session_state.pagina = "Sobre"
        st.rerun()

    st.markdown("---")

    if st.button("🚪 Excluir Seção"):
        st.session_state.clear()
        st.switch_page("pages/cadastro.py")

# ========== CONTEÚDO PRINCIPAL ==========
st.title("🛒 SmartList")
st.write("Gerencie sua lista de compras, controle seu orçamento e visualize comparações mensais de gastos.")

# ========== PÁGINA HOME ==========
if st.session_state.pagina == "Home":
    st.header("🏠 Visão Geral")
    
    total_itens = len(st.session_state.itens)
    total_comprados = sum(1 for i in st.session_state.itens if i["status"] == "Comprado")
    total_pendentes = total_itens - total_comprados
    total_gasto = sum(i["quantidade"] * i["preco"] for i in st.session_state.itens)
    
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("📦 Itens", total_itens)
    c2.metric("✅ Comprados", total_comprados)
    c3.metric("⏳ Pendentes", total_pendentes)
    c4.metric("💰 Total Gasto", f"R$ {total_gasto:.2f}")
    
    st.markdown("---")
    
    orcamentos = st.session_state.get("orcamentos_mensais", {})
    st.subheader("🎯 Situação Geral dos Orçamentos")
    
    if not orcamentos:
        st.info("Informe orçamentos na aba Orçamento para ver a situação geral.")
    else:
        totais_por_mes = {}
        for item in st.session_state.itens:
            mes = item["mes"]
            total_item = item["quantidade"] * item["preco"]
            if mes not in totais_por_mes:
                totais_por_mes[mes] = 0
            totais_por_mes[mes] += total_item
        
        meses_fora = 0
        meses_dentro = 0
        meses_sem_orcamento = 0
        
        for mes, total in totais_por_mes.items():
            orcamento = orcamentos.get(mes, 0)
            if orcamento == 0:
                meses_sem_orcamento += 1
            elif total > orcamento:
                meses_fora += 1
            else:
                meses_dentro += 1
        
        if meses_fora == 0 and meses_sem_orcamento == 0:
            st.success("✅ Todos os meses estão dentro do orçamento!")
        else:
            if meses_fora == 1:
                st.warning("⚠️ Há 1 mês que está fora do orçamento.")
            elif meses_fora > 1:
                st.error(f"❌ Há {meses_fora} meses que estão fora do orçamento.")
            if meses_sem_orcamento == 1:
                st.info("ℹ️ Há 1 mês que não possui orçamento definido.")
            elif meses_sem_orcamento > 1:
                st.info(f"ℹ️ Há {meses_sem_orcamento} meses que não possuem orçamento definido.")

# ========== PÁGINA CADASTRO ==========
elif st.session_state.pagina == "Cadastro":
    st.header("➕ Cadastro / Edição de Itens")
    
    # Configuração do formulário
    editando = st.session_state.editando
    if editando is not None and 0 <= editando < len(st.session_state.itens):
        item_edit = st.session_state.itens[editando]
        nome_default = item_edit.get("nome", "")
        quantidade_default = item_edit.get("quantidade", 1)
        preco_default = item_edit.get("preco", 0.0)
        mes_default = item_edit.get("mes", "Janeiro")
        status_default = item_edit["status"]
        botao_texto = "💾 Salvar Alterações"
    else:
        nome_default = ""
        quantidade_default = 1
        preco_default = 0.0
        mes_default = "Janeiro"
        status_default = "Pendente"
        botao_texto = "➕ Adicionar Item"
    
    with st.form("form_itens", clear_on_submit=False):
        col1, col2, col3 = st.columns([3, 1, 1])
        with col1:
            nome = st.text_input("📝 Nome do item", value=nome_default, placeholder="Ex: Arroz, Feijão...")
        with col2:
            quantidade = st.number_input("🔢 Quantidade", min_value=1, step=1, value=quantidade_default)
        with col3:
            preco = st.number_input("💰 Preço por unidade (R$)", min_value=0.0, step=0.10, format="%.2f", value=preco_default)
        
        mes = st.selectbox("📅 Mês da compra", MESES, index=MESES.index(mes_default) if mes_default in MESES else 0)
        status = st.selectbox(
            "📌 Status",
            ["Pendente", "Comprado"],
            index=0 if status_default == "Pendente" else 1
        )
        
        enviar = st.form_submit_button(botao_texto)
        
        if enviar:
            if not nome.strip():
                st.error("❌ O nome do item é obrigatório.")
            else:
                novo_item = {
                    "nome": nome.strip().capitalize(),
                    "quantidade": quantidade,
                    "preco": preco,
                    "mes": mes,
                    "status": status
                }
                
                if editando is not None:
                    st.session_state.itens[editando] = novo_item
                    st.success("✅ Item atualizado com sucesso!")
                    st.session_state.editando = None
                else:
                    st.session_state.itens.append(novo_item)
                    st.success("✅ Item adicionado com sucesso!")
                st.rerun()
    
    st.markdown("---")
    st.subheader("📋 Itens cadastrados (clique em opções para modificar)")
    
    if not st.session_state.itens:
        st.info("📝 Nenhum item cadastrado ainda.")
    else:
        for i, item in enumerate(st.session_state.itens):
            colA, colB, colC, colD, colE, colF = st.columns([3, 1, 1, 1, 1, 2])
            
            colA.markdown(f"**{item['nome']}**")
            colB.write(f"{item['quantidade']} un.")
            colC.write(f"R$ {item['preco']:.2f}")
            colD.write(item['mes'])
            colE.write(item['status'])
            
            # Gerenciar estado das opções
            if i not in st.session_state.opcoes_cadastro:
                st.session_state.opcoes_cadastro[i] = False
            
            botao_label = "❌ Ocultar" if st.session_state.opcoes_cadastro[i] else "⚙️ Opções"
            
            if colF.button(botao_label, key=f"toggle_cad_{i}", use_container_width=True):
                st.session_state.opcoes_cadastro[i] = not st.session_state.opcoes_cadastro[i]
                st.rerun()
            
            if st.session_state.opcoes_cadastro[i]:
                colX, colY = st.columns([1, 1])
                
                if colX.button("✏️ Editar", key=f"edit_cad_{i}", use_container_width=True):
                    st.session_state.editando = i
                    st.session_state.aviso_edicao = True
                    st.rerun()
                
                if colY.button("🗑️ Excluir", key=f"del_cad_{i}", use_container_width=True):
                    st.session_state.itens.pop(i)
                    st.warning("🗑️ Item excluído.")
                    st.rerun()
                
                if st.session_state.aviso_edicao:
                    st.warning("⬆️ Para editar este item, role a página para cima até o formulário.")
                    st.session_state.aviso_edicao = False
            
            st.markdown("---")

# ========== PÁGINA LISTAGEM ==========
elif st.session_state.pagina == "Listagem":
    st.header("📋 Listagem de Itens")
    
    # Botão de exportação
    if st.button("📤 Exportar para Excel"):
        if not st.session_state.itens:
            st.warning("Nenhum item para exportar.")
        else:
            nome_rel = st.session_state.get("usuario_nome", "Usuário")

            wb = Workbook()
            ws = wb.active
            ws.title = "Relatório"

            borda_fina = Border(
                left=Side(style="thin", color="000000"),
                right=Side(style="thin", color="000000"),
                top=Side(style="thin", color="000000"),
                bottom=Side(style="thin", color="000000")
            )

            cinza_escuro = PatternFill("solid", fgColor="5D6D7E")
            cinza_medio = PatternFill("solid", fgColor="85929E")
            cinza_claro = PatternFill("solid", fgColor="D6DBDF")
            verde_claro = PatternFill("solid", fgColor="D5F4E6")
            amarelo_claro = PatternFill("solid", fgColor="FEF9E7")

            ws.merge_cells("A1:G1")
            ws["A1"] = "RELATÓRIO GERAL - SMARTLIST"
            ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
            ws["A1"].alignment = Alignment(horizontal="center")
            ws["A1"].fill = cinza_escuro

            ws.merge_cells("A2:G2")
            ws["A2"] = f"Gerado em: {time.strftime('%d/%m/%Y %H:%M:%S')} | Usuário: {nome_rel}"
            ws["A2"].alignment = Alignment(horizontal="center")
            ws["A2"].font = Font(size=10, color="555555")

            ws.append([])

            cabecalhos = ["ID", "Nome", "Quantidade", "Preço", "Total", "Mês", "Status"]
            ws.append(cabecalhos)

            for col in range(1, 8):
                c = ws.cell(row=4, column=col)
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = cinza_medio
                c.border = borda_fina
                c.alignment = Alignment(horizontal="center")

            linha = 5
            total_geral = 0
            total_comprado = 0
            total_pendente = 0
            totais_mes = {}

            # CORREÇÃO: ID começa em 0
            for idx, item in enumerate(st.session_state.itens):
                total_item = item["quantidade"] * item["preco"]
                total_geral += total_item

                mes = item["mes"]
                totais_mes[mes] = totais_mes.get(mes, 0) + total_item

                if item["status"] == "Comprado":
                    total_comprado += total_item
                    cor_status = verde_claro
                else:
                    total_pendente += total_item
                    cor_status = amarelo_claro

                ws.append([
                    idx,  # CORREÇÃO: ID começa em 0
                    item["nome"],
                    item["quantidade"],
                    f"R$ {item['preco']:.2f}".replace('.', ','),
                    f"R$ {total_item:.2f}".replace('.', ','),
                    item["mes"],
                    item["status"]
                ])

                for col in range(1, 8):
                    c = ws.cell(row=linha, column=col)
                    c.border = borda_fina
                    if col in [1, 3, 6, 7]:
                        c.alignment = Alignment(horizontal="center")
                    elif col in [4, 5]:
                        c.alignment = Alignment(horizontal="right")
                    if col == 7:
                        c.fill = cor_status

                    if linha % 2 == 0:
                        if col != 7:
                            c.fill = PatternFill("solid", fgColor="F8F9F9")

                linha += 1

            linha_total = linha + 1

            for col in [1, 2, 6, 7]:
                c = ws.cell(row=linha_total, column=col)
                c.value = ""

            for col in [3, 4, 5]:
                c = ws.cell(row=linha_total, column=col)
                c.border = borda_fina
                c.fill = cinza_claro

            ws.merge_cells(f"C{linha_total}:E{linha_total}")
            cell_total = ws[f"C{linha_total}"]
            valor_formatado = f"R$ {total_geral:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
            cell_total.value = f"TOTAL GERAL: {valor_formatado}"
            cell_total.font = Font(bold=True, size=12)
            cell_total.alignment = Alignment(horizontal="center")
            cell_total.fill = cinza_claro

            cell_total.border = Border(
                left=Side(style="thin", color="000000"),
                right=Side(style="thin", color="000000"),
                top=Side(style="thin", color="000000"),
                bottom=Side(style="thin", color="000000")
            )

            linha_secoes = linha_total + 3

            ws[f"B{linha_secoes}"] = "RESUMO POR MÊS"
            ws[f"B{linha_secoes}"].font = Font(bold=True, color="FFFFFF", size=11)
            ws[f"B{linha_secoes}"].alignment = Alignment(horizontal="center")
            ws[f"B{linha_secoes}"].fill = cinza_escuro
            ws[f"B{linha_secoes}"].border = borda_fina
            ws[f"C{linha_secoes}"].fill = cinza_escuro
            ws[f"C{linha_secoes}"].border = borda_fina
            ws.merge_cells(f"B{linha_secoes}:C{linha_secoes}")

            linha_resumo_cab = linha_secoes + 1
            ws[f"B{linha_resumo_cab}"] = "Mês"
            ws[f"C{linha_resumo_cab}"] = "Total"
            
            for col in [2, 3]:
                c = ws.cell(row=linha_resumo_cab, column=col)
                c.font = Font(bold=True)
                c.fill = cinza_claro
                c.border = borda_fina
                c.alignment = Alignment(horizontal="center")

            linha_resumo = linha_resumo_cab + 1
            meses_ordenados = sorted(totais_mes.keys())
            
            for mes in meses_ordenados:
                valor = totais_mes[mes]
                ws[f"B{linha_resumo}"] = mes
                ws[f"C{linha_resumo}"] = f"R$ {valor:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
                
                ws[f"B{linha_resumo}"].border = borda_fina
                ws[f"C{linha_resumo}"].border = borda_fina
                ws[f"B{linha_resumo}"].alignment = Alignment(horizontal="center")
                ws[f"C{linha_resumo}"].alignment = Alignment(horizontal="right")

                if (linha_resumo - linha_resumo_cab) % 2 == 0:
                    ws[f"B{linha_resumo}"].fill = PatternFill("solid", fgColor="F8F9F9")
                    ws[f"C{linha_resumo}"].fill = PatternFill("solid", fgColor="F8F9F9")
                
                linha_resumo += 1

            ws[f"E{linha_secoes}"] = "INDICADORES"
            ws[f"E{linha_secoes}"].font = Font(bold=True, color="FFFFFF", size=11)
            ws[f"E{linha_secoes}"].alignment = Alignment(horizontal="center")
            ws[f"E{linha_secoes}"].fill = cinza_escuro
            ws[f"E{linha_secoes}"].border = borda_fina
            ws[f"F{linha_secoes}"].fill = cinza_escuro
            ws[f"F{linha_secoes}"].border = borda_fina
            ws.merge_cells(f"E{linha_secoes}:F{linha_secoes}")

            linha_ind = linha_secoes + 1
            
            indicadores = [
                ("Total Comprado", f"R$ {total_comprado:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')),
                ("Total Pendente", f"R$ {total_pendente:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')),
                ("% Concluído", f"{(total_comprado / total_geral * 100):.1f}%" if total_geral > 0 else "0.0%"),
                ("Média por Item", f"R$ {(total_geral/len(st.session_state.itens)):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.') if st.session_state.itens else "R$ 0,00"),
                ("Total de Itens", str(len(st.session_state.itens)))
            ]

            for i, (nome, valor) in enumerate(indicadores, 1):
                ws[f"E{linha_ind}"] = nome
                ws[f"F{linha_ind}"] = valor
                
                ws[f"E{linha_ind}"].border = borda_fina
                ws[f"F{linha_ind}"].border = borda_fina
                ws[f"E{linha_ind}"].alignment = Alignment(horizontal="left")
                ws[f"F{linha_ind}"].alignment = Alignment(horizontal="right")

                if "%" in valor:
                    percent_val = float(valor.replace('%', ''))
                    if percent_val >= 75:
                        ws[f"F{linha_ind}"].font = Font(color="2E7D32")
                    elif percent_val >= 50:
                        ws[f"F{linha_ind}"].font = Font(color="F57C00")
                    else:
                        ws[f"F{linha_ind}"].font = Font(color="C62828")

                if i % 2 == 0:
                    ws[f"E{linha_ind}"].fill = PatternFill("solid", fgColor="F8F9F9")
                    ws[f"F{linha_ind}"].fill = PatternFill("solid", fgColor="F8F9F9")
                
                linha_ind += 1

            ws.column_dimensions['A'].width = 8
            ws.column_dimensions['B'].width = 25
            ws.column_dimensions['C'].width = 15
            ws.column_dimensions['D'].width = 18
            ws.column_dimensions['E'].width = 18
            ws.column_dimensions['F'].width = 15
            ws.column_dimensions['G'].width = 15

            ws.freeze_panes = "A5"

            ws.auto_filter.ref = f"A4:G{linha}"

            caminho = f"relatorio_smartlist_{int(time.time())}.xlsx"
            wb.save(caminho)

            with open(caminho, "rb") as f:
                st.download_button(
                    "✅ Baixar Excel",
                    data=f,
                    file_name=caminho,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    help="Clique para baixar o relatório em Excel"
                )

    col1, col2, col3 = st.columns(3)
    with col1:
        search = st.text_input("🔎 Buscar item (nome)")
    with col2:
        status_filter = st.selectbox("📌 Filtrar status", ["Todos", "Pendente", "Comprado"])
    with col3:
        sort_by = st.selectbox("↕️ Ordenar por", ["Padrão", "Nome A-Z", "Preço (maior primeiro)", "Mês"])

    if not st.session_state.itens:
        st.info("Nenhum item cadastrado.")
    else:
        dados = []

        for idx, item in enumerate(st.session_state.itens):
            dados.append({
                "ID": idx,
                "Nome": item["nome"],
                "Quantidade": item["quantidade"],
                "Preço (R$)": item["preco"],
                "Total (R$)": item["quantidade"] * item["preco"],
                "Mês": item["mes"],
                "Status": item["status"]
            })

        if search:
            dados = [d for d in dados if search.lower() in d["Nome"].lower()]

        if status_filter != "Todos":
            dados = [d for d in dados if d["Status"] == status_filter]

        if sort_by == "Nome A-Z":
            dados.sort(key=lambda x: x["Nome"].lower())
        elif sort_by == "Preço (maior primeiro)":
            dados.sort(key=lambda x: x["Total (R$)"], reverse=True)
        elif sort_by == "Mês":
            dados.sort(key=lambda x: MESES.index(x["Mês"]))

        if not dados:
            st.info("Nenhum item encontrado com os filtros.")
        else:
            st.subheader("🗂️ Itens na Tabela")

            if search and not st.session_state.mostrar_tabela:
                st.warning("🔽 Expanda a tabela para visualizar os resultados da busca.")

            botao_tabela = "🔽 Mostrar Tabela" if not st.session_state.mostrar_tabela else "🔼 Ocultar Tabela"
            if st.button(botao_tabela):
                st.session_state.mostrar_tabela = not st.session_state.mostrar_tabela
                st.rerun()

            if st.session_state.mostrar_tabela:
                st.dataframe(
                    dados,
                    use_container_width=True,
                    hide_index=True
                )

            st.markdown("---")
            st.subheader("⚙️ Ações no Item")

            nomes = [f"ID {d['ID']} - {d['Nome']} ({d['Mês']})" for d in dados]
            item_escolhido = st.selectbox("Selecione um item", nomes)

            item_sel = dados[nomes.index(item_escolhido)]
            idx_real = item_sel["ID"]

            status_atual = st.session_state.itens[idx_real]["status"]
            texto_status = "✅ Marcar como Comprado" if status_atual == "Pendente" else "↩️ Marcar como Pendente"

            col1, col2, col3 = st.columns(3)

            if col1.button(texto_status, use_container_width=True):
                if status_atual == "Comprado":
                    st.session_state.itens[idx_real]["status"] = "Pendente"
                else:
                    st.session_state.itens[idx_real]["status"] = "Comprado"
                st.rerun()

            if col2.button("✏️ Editar", use_container_width=True):
                st.session_state.editando = idx_real
                st.session_state.pagina = "Cadastro"
                st.rerun()

            if col3.button("🗑️ Excluir", use_container_width=True):
                st.session_state.itens.pop(idx_real)
                st.warning("🗑️ Item excluído.")
                st.rerun()

# ========== PÁGINA ORÇAMENTO ==========
elif st.session_state.pagina == "Orcamento":
    st.header("💰 Orçamento Mensal Inteligente")

    if not st.session_state.itens:
        st.info("📝 Nenhum item registrado.")
        st.stop()

    totais_por_mes = {}
    for item in st.session_state.itens:
        mes = item["mes"]
        total_item = item["quantidade"] * item["preco"]

        if mes not in totais_por_mes:
            totais_por_mes[mes] = 0

        totais_por_mes[mes] += total_item

    st.subheader("📅 Ajustar Orçamento por Mês")

    with st.expander("🗂️ Clique para abrir os meses e ajustar os orçamentos"):
        for mes, total in totais_por_mes.items():
            valor_atual = st.session_state.orcamentos_mensais.get(mes, 0.0)

            col1, col2, col3 = st.columns([2,2,2])

            col1.markdown(f"**{mes}**")
            col2.markdown(f"Gasto: R$ {total:.2f}")

            novo_valor = col3.number_input(
                f"Orçamento {mes}",
                min_value=0.0,
                step=10.0,
                value=float(valor_atual),
                key=f"orc_{mes}"
            )

            if novo_valor != valor_atual:
                st.session_state.orcamentos_mensais[mes] = novo_valor
                st.toast("✅ Orçamento atualizado!", icon="✅")


    st.markdown("---")
    st.subheader("📊 Situação Financeira por Mês")

    for mes, total in totais_por_mes.items():
        orcamento = st.session_state.orcamentos_mensais.get(mes, 0)

        colA, colB, colC = st.columns(3)

        colA.metric(f"📅 {mes}", f"R$ {total:.2f}")
        colB.metric("🎯 Orçamento", f"R$ {orcamento:.2f}")

        if orcamento > 0:
            diferenca = orcamento - total
            percentual = (total / orcamento) * 100

            if percentual > 100:
                colC.error(f"❌ Estourado em R$ {-diferenca:.2f}")
            elif percentual > 80:
                colC.warning(f"⚠️ {percentual:.1f}% usado")
            else:
                colC.success(f"✅ {percentual:.1f}% usado")
        else:
            colC.info("Defina um orçamento")

        with st.expander(f"📋 Ver itens de {mes}"):

            tabela = []

            for item in st.session_state.itens:
                if item["mes"] == mes:
                    total_item = item["quantidade"] * item["preco"]

                    tabela.append({
                        "Item": item["nome"],
                        "Qtd": item["quantidade"],
                        "Preço (R$)": f"{item['preco']:.2f}",
                        "Total (R$)": f"{total_item:.2f}",
                        "Status": item["status"]
                    })

            if tabela:
                st.dataframe(tabela, use_container_width=True, hide_index=True)
            else:
                st.info("Nenhum item neste mês.")

    st.markdown("---")

# ========== PÁGINA HISTÓRICO ==========
elif st.session_state.pagina == "Historico":
    st.header("📅 Histórico Mensal (Automático)")

    if not st.session_state.itens:
        st.info("Nenhum item cadastrado ainda.")
    else:
        historico = {}

        for item in st.session_state.itens:
            mes = item["mes"]
            total = item["quantidade"] * item["preco"]

            if mes not in historico:
                historico[mes] = 0

            historico[mes] += total

        st.subheader("📊 Total de Gastos por Mês")

        if "mostrar_tabela_historico" not in st.session_state:
            st.session_state.mostrar_tabela_historico = False

        botao = "📂 Mostrar Histórico" if not st.session_state.mostrar_tabela_historico else "📁 Ocultar Tabela"

        if st.button(botao):
            st.session_state.mostrar_tabela_historico = not st.session_state.mostrar_tabela_historico
            st.rerun()

        if st.session_state.mostrar_tabela_historico:
            tabela = []

            for mes, valor in historico.items():
                tabela.append({
                    "Mês": mes,
                    "Total Gasto (R$)": f"{valor:.2f}"
                })

            st.dataframe(tabela, use_container_width=True, hide_index=True)

        media = sum(historico.values()) / len(historico)
        st.metric("💰 Média Mensal", f"R$ {media:.2f}")

        if len(historico) >= 2:
            st.subheader("📈 Comparação Entre Meses")

            meses = list(historico.keys())
            col1, col2 = st.columns(2)

            with col1:
                mes1 = st.selectbox("Primeiro mês", meses)
            with col2:
                mes2 = st.selectbox("Segundo mês", meses, index=1)

            if mes1 != mes2:
                v1 = historico[mes1]
                v2 = historico[mes2]
                dif = v2 - v1
                perc = (dif / v1 * 100) if v1 > 0 else 0

                st.markdown("---")
                st.metric(mes1, f"R$ {v1:.2f}")
                st.metric(mes2, f"R$ {v2:.2f}", delta=f"R$ {dif:.2f} ({perc:+.1f}%)")

                if dif > 0:
                    st.error(f"{mes2} teve aumento de R$ {dif:.2f}")
                elif dif < 0:
                    st.success(f"{mes2} teve redução de R$ {-dif:.2f}")
                else:
                    st.info("Os dois meses tiveram o mesmo gasto.")
            else:
                st.warning("⚠️ Por favor, escolha dois meses diferentes para comparar.")

# ========== PÁGINA SOBRE ==========
elif st.session_state.pagina == "Sobre":
    st.header("ℹ️ Sobre o SmartList")
    st.markdown("""
    O **SmartList** é um sistema inteligente desenvolvido para facilitar o controle de compras, 
    acompanhar gastos mensais e melhorar sua organização financeira no dia a dia.
    """)
    st.markdown("---")

    col1, col2, col3 = st.columns(3)
    with col1:
        st.markdown("### 📌 Nome do Projeto")
        st.success("SmartList – Gerenciador Inteligente de Compras")
    with col2:
        st.markdown("### 📘 Tema")
        st.info("Organização de compras, orçamento e acompanhamento mensal de gastos.")
    with col3:
        st.markdown("### 🎯 Objetivo Geral")
        st.warning("Facilitar a gestão financeira cotidiana.")

    st.markdown("---")
    st.markdown("### 🎯 Objetivos Específicos")
    st.markdown("""
    - Registrar itens  
    - Calcular total  
    - Exibir alertas de orçamento  
    - Registrar gastos mensais  
    - Comparar meses  
    - Melhorar educação financeira  
    """)

    st.markdown("---")
    colA, colB = st.columns(2)
    with colA:
        st.markdown("### 🛠 Tecnologias Utilizadas")
        st.info("""
        - Python  
        - Streamlit  
        - Excel  
        """)
    with colB:
        st.markdown("### 👥 Integrantes do Grupo")
        st.success("""
        - Marcelo de Sá V.
        - Eliada de Araújo L.
        - Antônio Wellincleiton Silva L.
        """)
    st.markdown("---")
    st.markdown("### 🧩 Funcionalidades")
    colX, colY = st.columns(2)
    with colX:
        st.markdown("""
        #### 📋 Lista de Compras  
        - Adição, edição e exclusão  
        - Marcar como comprado  
        - Cálculo automático por item  
        """)
    with colY:
        st.markdown("""
        #### 📅 Histórico Mensal  
        - Registro por mês  
        - Comparação entre meses  
        """)
    st.markdown("---")
    st.info("💡 Use o SmartList para desenvolver bons hábitos financeiros.")
    st.caption("© 2025 SmartList")